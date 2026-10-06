"""The Daily list's department filter, and pickers searchable by Employee ID
(2026-10-06)."""

import datetime
from decimal import Decimal

from django.urls import reverse
from django.utils import timezone

from attendance import tests_live
from attendance.services import recalculate
from common.tenant import use_company
from employees.services import create_employee
from organization.catalogue import adopt_department, adopt_designation

UTC = datetime.timezone.utc


class DepartmentFilterTests(tests_live.LiveTestCase):
    def setUp(self):
        super().setUp()
        with use_company(self.company):
            self.sales = adopt_department(self.branch, "SAL", "Sales")
            seller = adopt_designation(self.sales, "SLR", "Seller")
        self.karim = create_employee(
            company=self.company, first_name="Karim", employee_code="S7", branch=self.branch,
            department=self.sales, designation=seller,
            effective_from=datetime.datetime(2026, 1, 1, tzinfo=UTC),
            pay_basis="monthly", base_rate=Decimal("20000"))["employee"]
        self.day = timezone.localdate() - datetime.timedelta(days=3)
        recalculate(self.company.pk, start=self.day, end=self.day)
        self.client.force_login(self.admin)

    def names(self, **query):
        page = self.client.get(reverse("attendance:attendance_list"),
                               {"on": self.day.isoformat(), **query})
        self.assertEqual(page.status_code, 200)
        return {record.employee.first_name for record in page.context["page"].object_list}

    def test_the_list_narrows_to_a_department(self):
        self.assertEqual(self.names(), {"Rahim", "Karim"})
        self.assertEqual(self.names(department=self.sales.pk), {"Karim"})

    def test_the_filter_offers_the_departments(self):
        page = self.client.get(reverse("attendance:attendance_list"))
        self.assertContains(page, 'name="department"')
        self.assertContains(page, "All departments")
        self.assertContains(page, f'value="{self.sales.pk}"')

    def test_the_download_says_it(self):
        got = self.client.get(reverse("attendance:attendance_list"),
                              {"on": self.day.isoformat(), "department": self.sales.pk,
                               "format": "xlsx"})
        self.assertEqual(got.status_code, 200)
        from io import BytesIO

        from openpyxl import load_workbook

        text = " ".join(str(cell.value) for row in load_workbook(BytesIO(got.content)).active
                        .iter_rows() for cell in row if cell.value is not None)
        self.assertIn("Department: Sales", text)
        self.assertIn("Karim", text)
        self.assertNotIn("Rahim", text)


class EmployeeIdPickerTests(tests_live.LiveTestCase):
    def test_the_calendar_picker_shows_the_id_and_always_searches(self):
        self.client.force_login(self.admin)
        page = self.client.get(reverse("attendance:attendance_calendar"))
        self.assertContains(page, "E1 · Rahim")
        self.assertContains(page, 'id="employee" name="employee" data-search="always"')

    def test_the_daily_list_picker_shows_the_id(self):
        self.client.force_login(self.admin)
        self.assertContains(self.client.get(reverse("attendance:attendance_list")),
                            "E1 · Rahim")
