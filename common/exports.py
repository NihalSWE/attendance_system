"""Excel and PDF downloads of what a page is showing.

The pages build their own rows - with their own scoping, filters, search and
sort, and their own rules about what a viewer may see - and hand them here.
This module only writes the file, and holds the rules every download shares:

- **It says what it is.** A title, the company, and a line of the filters and
  the row count at the top of every file - a header block in Excel, a subtitle
  in PDF - so a printed sheet can be read without the screen it came from.
- **It has a ceiling.** Above ``MAX_ROWS[format]`` the download is refused with
  a message naming the ceiling and asking for a narrower filter. Nothing is
  truncated silently, and one click cannot hold a worker for a minute.
- **It is recorded.** Each download is a company event (``export.downloaded``)
  with who, which page, which format, which filters and how many rows.
- **It is named for what it holds:** ``employees-headoffice-2026-09-21.xlsx``.

PDF uses reportlab (pure Python; no system libraries on the server) with one
bundled font for the whole document, **Hind Siliguri** (common/fonts/, SIL Open
Font License): it has Bangla *and* Latin, so a Bangla name and an English one
print from the same font. (Noto Sans Bengali was the first choice but carries
no Latin letters, digits or punctuation - everything English would have been
boxes.) Bangla needs shaping - conjuncts, and vowel signs drawn before their
consonant - which reportlab does through ``uharfbuzz``; it is in
requirements.txt, and ``_register_fonts`` refuses to run without it rather
than print names in the wrong letter order.
"""

import datetime
import io
import re
from pathlib import Path

from django.http import HttpResponse

from auditlog.services import record_company_event

XLSX = "xlsx"
PDF = "pdf"
FORMATS = (XLSX, PDF)

#: The most rows one download may hold. Excel is cheap to write; a PDF of this
#: many rows is already ~40 landscape pages and a few seconds of work.
MAX_ROWS = {XLSX: 10_000, PDF: 1_500}

CONTENT_TYPES = {
    XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    PDF: "application/pdf",
}


class TooManyRows(Exception):
    """The download would exceed the ceiling; ``str()`` is the reader's message."""


def check_size(count, fmt, noun="rows"):
    ceiling = MAX_ROWS[fmt]
    if count > ceiling:
        label = "Excel" if fmt == XLSX else "PDF"
        raise TooManyRows(
            f"That is {count:,} {noun}. A {label} download holds at most "
            f"{ceiling:,}. Narrow the filter - one branch, a shorter date range, "
            "or a search - and download again."
        )


def slug(text):
    """``Head Office`` -> ``headoffice``: letters and digits only, for filenames."""
    return re.sub(r"[^a-z0-9]", "", (text or "").casefold()) or "company"


def filename(page, scope, period, fmt):
    return f"{page}-{slug(scope)}-{period}.{fmt}"


def record(*, actor, membership, page, fmt, filters, count):
    """The audit line for a download: salary and attendance leaving the system."""
    record_company_event(
        actor=actor, membership=membership, company=membership.company,
        action="export.downloaded", obj=membership.company,
        after={"page": page, "format": fmt, "filters": filters, "rows": count},
    )


def response(content, fmt, name):
    reply = HttpResponse(content, content_type=CONTENT_TYPES[fmt])
    reply["Content-Disposition"] = f'attachment; filename="{name}"'
    return reply


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        return value.strftime("%d %b %Y %H:%M")
    if isinstance(value, datetime.date):
        return value.strftime("%d %b %Y")
    return str(value)


# --- Excel -----------------------------------------------------------------


def table_xlsx(*, title, lines, headers, rows, numeric=(), widths=None):
    """One sheet: the title and the filter lines, a blank row, then the table.

    ``rows`` are lists of plain values; numbers stay numbers (``numeric`` names
    the column indexes to right-align and format with thousands separators),
    dates stay dates, so the sheet can be summed and sorted.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = title[:31]
    sheet.append([title])
    sheet["A1"].font = Font(bold=True, size=13)
    for line in lines:
        sheet.append([line])
    sheet.append([])
    header_row = sheet.max_row + 1
    sheet.append(list(headers))
    fill = PatternFill("solid", fgColor="E8ECF1")
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
        cell.fill = fill
    for row in rows:
        sheet.append([value if value is not None else "" for value in row])
    for index in numeric:
        letter = get_column_letter(index + 1)
        for (cell,) in sheet.iter_rows(min_row=header_row + 1, min_col=index + 1,
                                       max_col=index + 1):
            cell.alignment = Alignment(horizontal="right")
            if isinstance(cell.value, (int, float)) or hasattr(cell.value, "as_tuple"):
                cell.number_format = "#,##0.00"
        sheet[f"{letter}{header_row}"].alignment = Alignment(horizontal="right")
    for index, heading in enumerate(headers, start=1):
        width = (widths or {}).get(index - 1) or max(12, min(40, len(str(heading)) + 6))
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = f"A{header_row + 1}"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# --- PDF -------------------------------------------------------------------


#: The one font family every PDF uses (see the module docstring).
FONT = "HindSiliguri"
FONT_BOLD = "HindSiliguri-Bold"
FONT_DIR = Path(__file__).resolve().parent / "fonts"


def _register_fonts():
    """Register the bundled font with reportlab, once per process."""
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if FONT in pdfmetrics.getRegisteredFontNames():
        return
    try:
        import uharfbuzz  # noqa: F401 - reportlab shapes Bangla only with it
    except ModuleNotFoundError:  # pragma: no cover - listed in requirements.txt
        raise RuntimeError(
            "uharfbuzz is not installed: Bangla would print in the wrong letter "
            "order. Run pip install -r requirements.txt."
        )
    pdfmetrics.registerFont(TTFont(FONT, str(FONT_DIR / "HindSiliguri-Regular.ttf")))
    pdfmetrics.registerFont(TTFont(FONT_BOLD, str(FONT_DIR / "HindSiliguri-Bold.ttf")))
    # So <b> inside a paragraph picks the bold face of the same family.
    addMapping(FONT, 0, 0, FONT)
    addMapping(FONT, 1, 0, FONT_BOLD)
    addMapping(FONT, 0, 1, FONT)
    addMapping(FONT, 1, 1, FONT_BOLD)


def _styles():
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet

    _register_fonts()
    base = getSampleStyleSheet()
    # shaping=1 on every style: reportlab defaults it OFF, and unshaped Bangla
    # prints its letters in the wrong order (a vowel sign after its consonant
    # instead of before). Bangla sits taller than Latin, so the leading is a
    # little looser than Helvetica needed.
    return {
        "title": ParagraphStyle("t", parent=base["Title"], fontName=FONT_BOLD, fontSize=14,
                                leading=19, alignment=0, spaceAfter=2, shaping=1),
        "sub": ParagraphStyle("s", parent=base["Normal"], fontName=FONT, fontSize=8,
                              leading=11, textColor=colors.HexColor("#4a5a6b"), shaping=1),
        "cell": ParagraphStyle("c", parent=base["Normal"], fontName=FONT, fontSize=7.5,
                               leading=10, shaping=1),
        "head": ParagraphStyle("h", parent=base["Normal"], fontName=FONT_BOLD, fontSize=7.5,
                               leading=10, shaping=1),
    }


def _document(buffer, title, landscape_page=True):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate

    size = landscape(A4) if landscape_page else A4

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(FONT, 7)
        canvas.drawRightString(size[0] - 12 * mm, 8 * mm, f"{title} · page {doc.page}",
                               shaping=True)
        canvas.restoreState()

    # initialFontName: reportlab opens every page with an (empty) text block in
    # its default Helvetica; starting on our font keeps Helvetica out entirely.
    doc = SimpleDocTemplate(buffer, pagesize=size, title=title, initialFontName=FONT,
                            leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm)
    return doc, footer


def _escape(text):
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def table_pdf(*, title, lines, headers, rows, numeric=(), widths=None):
    """A landscape A4 table: title, filter lines, then the rows.

    The header repeats on every page and long cells wrap. ``widths`` are
    relative column weights; the table fills the page width.
    """
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

    styles = _styles()
    buffer = io.BytesIO()
    doc, footer = _document(buffer, title)
    story = [Paragraph(_escape(title), styles["title"])]
    story += [Paragraph(_escape(line), styles["sub"]) for line in lines]
    story.append(Spacer(1, 6))

    right = styles["cell"].clone("r", alignment=2)
    head_right = styles["head"].clone("hr", alignment=2)
    data = [[Paragraph(_escape(str(h)), head_right if i in numeric else styles["head"])
             for i, h in enumerate(headers)]]
    for row in rows:
        data.append([Paragraph(_escape(_cell_text(value)), right if i in numeric else styles["cell"])
                     for i, value in enumerate(row)])

    weights = [(widths or {}).get(i, 1) for i in range(len(headers))]
    total = sum(weights)
    col_widths = [doc.width * w / total for w in weights]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), FONT),    # not the table default, Helvetica
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8ecf1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f8fa")]),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#9aa7b5")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story.append(table)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def grid_pdf(*, title, lines, blocks):
    """Calendar grids, one block after another.

    ``blocks`` is a list of ``(heading, header_row, week_rows)``: a heading
    line (a person's name, say), the weekday names, and one list per week of
    cell texts - a cell may hold several lines separated by newlines.
    """
    from reportlab.lib import colors
    from reportlab.platypus import KeepTogether, Paragraph, Spacer, Table, TableStyle

    styles = _styles()
    buffer = io.BytesIO()
    doc, footer = _document(buffer, title)
    story = [Paragraph(_escape(title), styles["title"])]
    story += [Paragraph(_escape(line), styles["sub"]) for line in lines]
    story.append(Spacer(1, 6))
    cell = styles["cell"]
    for heading, header_row, weeks in blocks:
        data = [[Paragraph(f"<b>{_escape(h)}</b>", cell) for h in header_row]]
        for week in weeks:
            data.append([Paragraph(_escape(text).replace("\n", "<br/>"), cell)
                         for text in week])
        grid = Table(data, colWidths=[doc.width / len(header_row)] * len(header_row),
                     repeatRows=1)
        grid.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), FONT),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c3ccd6")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8ecf1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 1), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 1), (-1, -1), 6),
        ]))
        part = []
        if heading:
            part.append(Paragraph(f"<b>{_escape(heading)}</b>", styles["sub"]))
            part.append(Spacer(1, 3))
        part += [grid, Spacer(1, 10)]
        story.append(KeepTogether(part))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def document_pdf(*, title, lines, sections, summary=(), note=""):
    """A portrait document: title, its lines, then headed tables.

    For one person's paper - a payslip today - rather than a list of rows:
    ``sections`` is ``(heading, headers, rows, numeric)`` and headers may be
    empty for a plain label/value block; ``summary`` is the ``(label, value)``
    pairs of the box at the end. Same fonts, page furniture and footer as
    every other PDF here, so Bangla works the same way.
    """
    from reportlab.lib import colors
    from reportlab.platypus import KeepTogether, Paragraph, Spacer, Table, TableStyle

    styles = _styles()
    buffer = io.BytesIO()
    doc, footer = _document(buffer, title, landscape_page=False)
    story = [Paragraph(_escape(title), styles["title"])]
    story += [Paragraph(_escape(line), styles["sub"]) for line in lines]
    story.append(Spacer(1, 8))

    right = styles["cell"].clone("r", alignment=2)
    head_right = styles["head"].clone("hr", alignment=2)

    def cells(values, numeric, bold=False):
        style = styles["head"] if bold else styles["cell"]
        style_right = head_right if bold else right
        return [Paragraph(_escape(_cell_text(value)),
                          style_right if index in numeric else style)
                for index, value in enumerate(values)]

    for heading, headers, rows, numeric in sections:
        if not rows:
            continue
        part = []
        if heading:
            part.append(Paragraph(f"<b>{_escape(heading)}</b>", styles["sub"]))
            part.append(Spacer(1, 3))
        data = [cells(headers, numeric, bold=True)] if headers else []
        data += [cells(row, numeric) for row in rows]
        table = Table(data, colWidths=[doc.width / len(data[0])] * len(data[0]),
                      repeatRows=1 if headers else 0)
        style = [
            ("FONTNAME", (0, 0), (-1, -1), FONT),   # not the table default, Helvetica
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ("LINEBELOW", (0, 0), (-1, -2), 0.3, colors.HexColor("#dfe4ea")),
        ]
        if headers:
            style += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8ecf1")),
                      ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#9aa7b5"))]
        table.setStyle(TableStyle(style))
        part += [table, Spacer(1, 10)]
        story.append(KeepTogether(part))

    if summary:
        data = [[Paragraph(f"<b>{_escape(str(label))}</b>", styles["cell"]),
                 Paragraph(f"<b>{_escape(_cell_text(value))}</b>", right)]
                for label, value in summary]
        box = Table(data, colWidths=[doc.width * 0.7, doc.width * 0.3])
        box.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), FONT),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f6f8fa")),
            ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#9aa7b5")),
            ("LINEBELOW", (0, 0), (-1, -2), 0.3, colors.HexColor("#dfe4ea")),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(box)
    if note:
        story += [Spacer(1, 8), Paragraph(_escape(note), styles["sub"])]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
