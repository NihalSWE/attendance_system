"""Bulk employee import: Employee ID and Name in the file, the branch on the form.

TwoBranchCase: Head Office (the company's default branch, where Manny is the
branch manager) and Chittagong. The company picks the branch, starting on the
default; a branch manager is locked to their own. Everyone lands in the
branch's Unassigned department with no salary, exactly as a device import
does, and HR fills in the rest on Edit employee.
"""

import csv
import datetime
import io
from zoneinfo import ZoneInfo

from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from auditlog.models import AuditLog
from common.tenant import use_company
from devices.services.mapping import UNASSIGNED_CODE
from employees.models import Employee, EmployeeAssignment, EmployeeCompensation
from leaves.tests_branch_access import TwoBranchCase
from organization import import_services
from organization.models import Branch, Department


class ImportCase(TwoBranchCase):
    def setUp(self):
        super().setUp()
        self.url = reverse("organization:employee_import")
        self.confirm_url = reverse("organization:employee_import_confirm")

    def csv_file(self, rows, *, name="people.csv", headings=("Employee ID", "Name")):
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(headings)
        for row in rows:
            writer.writerow(row)
        return SimpleUploadedFile(name, buffer.getvalue().encode("utf-8"),
                                  content_type="text/csv")

    def upload(self, rows, *, user=None, branch=None, **file_options):
        self.client.force_login(user or self.admin)
        payload = {"upload": self.csv_file(rows, **file_options)}
        if branch is not None:
            payload["branch"] = branch.pk
        elif user is None or user == self.admin:
            payload["branch"] = self.branch.pk      # what the dropdown starts on
        return self.client.post(self.url, payload)

    def confirm(self):
        return self.client.post(self.confirm_url, follow=True)

    def count(self):
        with use_company(self.company):
            return Employee.objects.count()

    def placement(self, code):
        with use_company(self.company):
            return EmployeeAssignment.objects.select_related(
                "branch", "department", "designation", "employee").get(employee_code=code)

    def upload_error(self, page):
        return " ".join(str(m) for m in page.context["form"].errors.get("upload", []))


class DemoFileTests(ImportCase):
    def test_the_demo_file_downloads_in_the_format_it_asks_for(self):
        self.client.force_login(self.admin)
        page = self.client.get(reverse("organization:employee_import_demo"))
        self.assertEqual(page.status_code, 200)
        self.assertIn("text/csv", page["Content-Type"])
        self.assertIn("employee-import-demo.csv", page["Content-Disposition"])
        text = page.content.decode("utf-8-sig")
        self.assertEqual(text.splitlines()[0], "Employee ID,Name")
        # What the demo writes is what the import reads.
        rows = import_services.read_file(SimpleUploadedFile("demo.csv", page.content))
        self.assertEqual(len(rows), len(import_services.DEMO_ROWS))
        self.assertEqual((rows[0]["employee_id"], rows[0]["name"]), ("445961", "Sajal Ahmed"))

    def test_the_excel_demo_file_round_trips_through_the_parser(self):
        self.client.force_login(self.admin)
        page = self.client.get(reverse("organization:employee_import_demo") + "?format=xlsx")
        self.assertEqual(page.status_code, 200)
        self.assertIn("spreadsheetml", page["Content-Type"])
        self.assertIn("employee-import-demo.xlsx", page["Content-Disposition"])
        rows = import_services.read_file(SimpleUploadedFile("demo.xlsx", page.content))
        # Same people as the CSV one, in the same order, read the same way.
        csv_rows = import_services.read_file(
            SimpleUploadedFile("demo.csv", import_services.demo_csv()))
        self.assertEqual([(r["employee_id"], r["name"]) for r in rows],
                         [(r["employee_id"], r["name"]) for r in csv_rows])
        self.assertEqual(len(rows), len(import_services.DEMO_ROWS))

    def test_the_page_offers_both_demo_files(self):
        self.client.force_login(self.admin)
        page = self.client.get(self.url)
        self.assertContains(page, "Download demo file (CSV)")
        self.assertContains(page, "Download demo file (Excel)")

    def test_a_branch_manager_can_download_it_too(self):
        self.client.force_login(self.manager)
        page = self.client.get(reverse("organization:employee_import_demo"))
        self.assertEqual(page.status_code, 200)


class ReadingTheFileTests(ImportCase):
    def test_a_missing_heading_is_named(self):
        page = self.upload([["445962"]], headings=("Employee ID",))
        self.assertIn("Missing: Name", self.upload_error(page))

    def test_other_spellings_of_the_headings_are_accepted(self):
        for headings in (("EMP-ID", "Name"), ("emp id", "NAME"), ("EMPID", "Full name")):
            with self.subTest(headings=headings):
                rows = import_services.read_file(
                    self.csv_file([["445962", "Ajay Kumar"]], headings=headings))
                self.assertEqual(rows[0]["employee_id"], "445962")

    def test_extra_columns_are_ignored(self):
        rows = import_services.read_file(self.csv_file(
            [["x", "445962", "Ajay Kumar"]], headings=("Notes", "Employee ID", "Name")))
        self.assertEqual((rows[0]["employee_id"], rows[0]["name"]), ("445962", "Ajay Kumar"))

    def test_an_empty_file_is_refused(self):
        self.client.force_login(self.admin)
        page = self.client.post(self.url, {
            "branch": self.branch.pk,
            "upload": SimpleUploadedFile("e.csv", b"", content_type="text/csv")})
        self.assertTrue(page.context["form"].errors)
        self.assertEqual(self.count(), 3)

    def test_headings_with_no_people_is_refused(self):
        self.assertIn("no people", self.upload_error(self.upload([])))

    def test_another_file_type_is_refused(self):
        self.client.force_login(self.admin)
        page = self.client.post(self.url, {
            "branch": self.branch.pk,
            "upload": SimpleUploadedFile("people.pdf", b"%PDF-1.4")})
        self.assertIn(".csv", self.upload_error(page))

    def test_too_many_rows_is_refused(self):
        original = import_services.MAX_ROWS
        import_services.MAX_ROWS = 2
        try:
            page = self.upload([[str(500000 + n), f"P {n}"] for n in range(3)])
        finally:
            import_services.MAX_ROWS = original
        self.assertIn("3 people", self.upload_error(page))

    def test_an_excel_file_is_read_too(self):
        from openpyxl import Workbook

        book = Workbook()
        book.active.append(["Employee ID", "Name"])
        book.active.append([445962, "Ajay Kumar"])        # a number, as Excel keeps it
        buffer = io.BytesIO()
        book.save(buffer)
        rows = import_services.read_file(SimpleUploadedFile("p.xlsx", buffer.getvalue()))
        self.assertEqual(rows[0]["employee_id"], "445962")   # not "445962.0"


class PreviewTests(ImportCase):
    def test_a_good_file_previews_and_writes_nothing(self):
        page = self.upload([["445962", "Ajay Kumar"], ["445963", "Dia Rahman"]])
        preview = page.context["preview"]
        self.assertEqual(preview["counts"], {"total": 2, "bad": 0, "good": 2})
        self.assertEqual(preview["branch"], self.branch)
        self.assertEqual(self.count(), 3)
        self.assertContains(page, "Import 2 employees")

    def test_every_bad_row_is_named_with_its_line(self):
        page = self.upload([
            ["44A962", "Letters In Id"],
            ["445963", "Good Person"],
            ["", "No Id"],
            ["445964", ""],
            ["445963", "Same Id Again"],
        ])
        preview = page.context["preview"]
        self.assertEqual([row["line"] for row in preview["bad"]], [2, 4, 5, 6])
        said = " ".join(e for row in preview["bad"] for e in row["errors"])
        for phrase in ("is not a number", "Employee ID is missing", "Name is missing",
                       "also on row 3"):
            self.assertIn(phrase, said)
        self.assertNotContains(page, "Import 1 employee")

    # Someone already in the software used to refuse the whole file, so a list
    # that had grown could never be uploaded again for its new people (Nihal,
    # 2026-09-24). They are now skipped and left unchanged: see ExistingPeopleTests.


class CompanyBranchTests(ImportCase):
    def test_the_dropdown_starts_on_the_default_branch_and_lists_every_branch(self):
        self.client.force_login(self.admin)
        form = self.client.get(self.url).context["form"]
        self.assertEqual(form.fields["branch"].initial, self.branch.pk)
        self.assertFalse(form.fields["branch"].disabled)
        with use_company(self.company):
            self.assertEqual(set(form.fields["branch"].queryset), {self.branch, self.unit})

    def test_leaving_the_default_puts_everyone_in_the_default_branch(self):
        self.upload([["445962", "Ajay Kumar"]])
        self.confirm()
        self.assertEqual(self.placement("445962").branch, self.branch)

    def test_choosing_another_branch_puts_everyone_there(self):
        page = self.upload([["445962", "Ajay Kumar"], ["445963", "Dia Rahman"]],
                           branch=self.unit)
        self.assertContains(page, "joining Chittagong")
        self.confirm()
        self.assertEqual(self.placement("445962").branch, self.unit)
        self.assertEqual(self.placement("445963").branch, self.unit)

    def test_clearing_the_dropdown_means_the_default_branch(self):
        self.client.force_login(self.admin)
        page = self.client.post(self.url, {
            "branch": "", "upload": self.csv_file([["445962", "Ajay Kumar"]])})
        self.assertEqual(page.context["preview"]["branch"], self.branch)
        self.confirm()
        self.assertEqual(self.placement("445962").branch, self.branch)

    def test_no_branch_at_all_means_the_default_in_the_service(self):
        self.assertEqual(import_services.check_branch(self.admin, self.company.pk), self.branch)


class BranchManagerTests(ImportCase):
    def test_the_branch_is_locked_to_their_own(self):
        self.client.force_login(self.manager)
        page = self.client.get(self.url)
        field = page.context["form"].fields["branch"]
        self.assertTrue(field.disabled)
        self.assertEqual(field.initial, self.branch.pk)
        self.assertContains(page, "You add people to your own branch.")

    def test_they_import_into_their_own_branch(self):
        page = self.upload([["445962", "Ajay Kumar"]], user=self.manager)
        self.assertEqual(page.context["preview"]["branch"], self.branch)
        self.confirm()
        self.assertEqual(self.placement("445962").branch, self.branch)

    def test_posting_another_branch_is_ignored(self):
        """The locked field keeps its own value whatever the request says."""
        page = self.upload([["445962", "Ajay Kumar"]], user=self.manager, branch=self.unit)
        self.assertEqual(page.context["preview"]["branch"], self.branch)
        self.confirm()
        self.assertEqual(self.placement("445962").branch, self.branch)

    def test_the_service_refuses_another_branch(self):
        with self.assertRaises(PermissionDenied):
            import_services.check_branch(self.manager, self.company.pk, self.unit.pk)
        with self.assertRaises(PermissionDenied):
            import_services.commit(actor=self.manager, company_id=self.company.pk,
                                   rows=[{"line": 2, "employee_id": "445962",
                                          "name": "Sneaky"}],
                                   branch_id=self.unit.pk)
        self.assertEqual(self.count(), 3)

    def test_a_manager_of_two_branches_chooses_between_those_two_only(self):
        with use_company(self.company):
            Branch.objects.create(company=self.company, code="SYL", name="Sylhet",
                                  timezone="Asia/Dhaka", country_code="BD")
        both = self.member("both@liv.test", "manager", branches=[self.branch, self.unit])
        self.client.force_login(both)
        field = self.client.get(self.url).context["form"].fields["branch"]
        self.assertFalse(field.disabled)
        with use_company(self.company):
            self.assertEqual(set(field.queryset), {self.branch, self.unit})


class ConfirmTests(ImportCase):
    def test_they_are_created_like_a_device_import(self):
        before = self.count()
        self.upload([["445962", "Ajay Kumar"], ["445963", "Md Fazle Rabbi"]])
        page = self.confirm()
        self.assertEqual(self.count(), before + 2)
        self.assertContains(page, "2 employees imported")

        placement = self.placement("445962")
        employee = placement.employee
        self.assertEqual((employee.first_name, employee.last_name), ("Ajay", "Kumar"))
        self.assertEqual(self.placement("445963").employee.full_name, "Md Fazle Rabbi")
        self.assertEqual(employee.employment_status, "active")
        self.assertTrue(employee.metadata["needs_hr_review"])
        # The branch's Unassigned department and designation, as a device import uses.
        self.assertEqual(placement.department.code, UNASSIGNED_CODE)
        self.assertEqual(placement.designation.code, UNASSIGNED_CODE)
        self.assertEqual(placement.department.branch_id, self.branch.pk)
        # No salary: payroll skips them by name until one is set.
        with use_company(self.company):
            self.assertFalse(EmployeeCompensation.objects.filter(employee=employee).exists())

    def test_they_start_at_midnight_today(self):
        self.upload([["445962", "Ajay Kumar"]])
        self.confirm()
        placement = self.placement("445962")
        zone = ZoneInfo(self.company.timezone or "UTC")
        local = placement.effective_from.astimezone(zone)
        self.assertEqual((local.hour, local.minute), (0, 0))
        self.assertEqual(local.date(), datetime.datetime.now(zone).date())
        self.assertEqual(placement.employee.joining_date, local.date())

    def test_setting_the_department_from_today_corrects_the_row(self):
        """The follow-up on Edit employee fixes the placement, adds no history."""
        from organization.employee_edit_services import change_placement

        self.upload([["445962", "Ajay Kumar"]])
        self.confirm()
        placement = self.placement("445962")
        change_placement(actor=self.admin, company_id=self.company.pk,
                         employee_id=placement.employee_id, values={
                             "branch": self.branch, "department": self.hq_department,
                             "designation": self.hq_designation,
                             "employee_code": "445962",
                             "effective_at": placement.effective_from, "reason": "",
                         })
        with use_company(self.company):
            rows = EmployeeAssignment.objects.filter(employee_id=placement.employee_id)
            self.assertEqual(rows.count(), 1)
            self.assertEqual(rows.get().department, self.hq_department)

    def test_a_second_import_reuses_the_same_unassigned_department(self):
        self.upload([["445962", "Ajay Kumar"]])
        self.confirm()
        self.upload([["445963", "Dia Rahman"]])
        self.confirm()
        with use_company(self.company):
            self.assertEqual(Department.objects.filter(
                branch=self.branch, code=UNASSIGNED_CODE).count(), 1)

    def test_one_bad_row_means_nothing_is_written(self):
        self.upload([["445962", "Good"], ["44A963", "Bad"]])
        page = self.confirm()
        self.assertEqual(self.count(), 3)
        self.assertContains(page, "can no longer be imported")

    def test_an_emp_id_taken_after_the_preview_rolls_everything_back(self):
        self.upload([["445962", "Ajay Kumar"], ["445963", "Dia Rahman"]])
        with use_company(self.company):
            EmployeeAssignment.objects.filter(employee_code="E1").update(employee_code="445963")
        page = self.confirm()
        self.assertEqual(self.count(), 3)
        self.assertContains(page, "can no longer be imported")

    def test_confirming_twice_imports_once(self):
        self.upload([["445962", "Ajay Kumar"]])
        self.confirm()
        before = self.count()
        self.assertContains(self.confirm(), "no longer waiting")
        self.assertEqual(self.count(), before)

    def test_the_import_is_audited_on_one_line(self):
        self.upload([["445962", "Ajay Kumar"], ["445963", "Dia Rahman"]])
        self.confirm()
        entry = AuditLog.objects.get(action="employees.imported")
        self.assertEqual(entry.actor_user_id, self.admin.pk)
        self.assertEqual(entry.after_data["count"], 2)
        self.assertEqual(entry.after_data["branch_id"], self.branch.pk)
        self.assertEqual(entry.after_data["employee_codes"], ["445962", "445963"])


class WhoMayImportTests(ImportCase):
    def test_an_employee_login_never_reaches_it(self):
        self.client.force_login(self.clerk_user)
        for url in (self.url, reverse("organization:employee_import_demo")):
            with self.subTest(url=url):
                page = self.client.get(url)
                self.assertEqual(page.status_code, 302)
                self.assertIn("/me/", page["Location"])
        self.assertIn(self.client.post(self.confirm_url).status_code, (302, 403))
        self.assertEqual(self.count(), 3)

    def test_the_services_refuse_them_too(self):
        rows = [{"line": 2, "employee_id": "445962", "name": "Sneaky"}]
        with self.assertRaises(PermissionDenied):
            import_services.check(self.clerk_user, self.company.pk, rows)
        with self.assertRaises(PermissionDenied):
            import_services.commit(actor=self.clerk_user, company_id=self.company.pk,
                                   rows=rows, branch_id=self.branch.pk)
        self.assertEqual(self.count(), 3)

    def test_the_employees_list_links_to_it(self):
        self.client.force_login(self.admin)
        self.assertContains(self.client.get(reverse("employee_list")), self.url)

    def test_the_page_shows_the_format_and_the_demo_link(self):
        self.client.force_login(self.admin)
        page = self.client.get(self.url)
        self.assertContains(page, "Employee ID")
        self.assertContains(page, "Download demo file (CSV)")
        self.assertIsNone(page.context["preview"])


# --- old Excel (.xls), and files that are not what they are named -----------

def biff2(rows):
    """A real Excel 97-2003 (BIFF) workbook, built here so no binary fixture
    and no customer file has to live in the repo. xlrd reads BIFF 2 onwards."""
    import struct

    def record(code, payload):
        return struct.pack("<HH", code, len(payload)) + payload

    out = [record(0x0009, struct.pack("<HH", 2, 0x0010))]      # BOF, worksheet
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if isinstance(value, (int, float)):
                out.append(record(0x0003, struct.pack("<HH3sd", r, c, b"\0\0\0", float(value))))
            else:
                text = str(value).encode("latin-1")
                out.append(record(0x0004, struct.pack("<HH3sB", r, c, b"\0\0\0", len(text)) + text))
    out.append(record(0x000A, b""))
    return b"".join(out)


class OldExcelTests(ImportCase):
    """A client's list arrived as Excel 97-2003 with an "ID" column (2026-09-23)."""

    def xls(self, rows, name="File_01.xls"):
        return SimpleUploadedFile(name, biff2(rows), content_type="application/vnd.ms-excel")

    def test_an_old_xls_is_read(self):
        rows = import_services.read_file(self.xls([
            ["SL", "ID", "Name"],
            [1, 830011, "Md. Hafizul Islam"],
            [2, 830032.0, "Biplob Baidya"],
        ]))
        # "ID" is the Employee ID; the serial column is ignored; a number
        # stays a number, not "830011.0".
        self.assertEqual([(r["employee_id"], r["name"]) for r in rows],
                         [("830011", "Md. Hafizul Islam"), ("830032", "Biplob Baidya")])

    def test_it_imports_through_the_page(self):
        self.client.force_login(self.admin)
        page = self.client.post(self.url, {
            "branch": self.branch.pk,
            "upload": self.xls([["ID", "Name"], [830011, "Md. Hafizul Islam"]]),
        })
        self.assertEqual(page.context["preview"]["counts"], {"total": 1, "bad": 0, "good": 1})
        self.confirm()
        self.assertEqual(self.placement("830011").employee.full_name, "Md. Hafizul Islam")

    def test_a_missing_heading_says_what_it_found(self):
        page = self.upload([], headings=("Serial", "Person"))
        message = self.upload_error(page)
        self.assertIn("Missing: Employee ID, Name", message)
        self.assertIn("Found: Serial, Person", message)

    def test_a_workbook_is_read_by_what_it_is_not_by_its_name(self):
        from openpyxl import Workbook

        book = Workbook()
        book.active.append(["ID", "Name"])
        book.active.append([830011, "Md. Hafizul Islam"])
        buffer = io.BytesIO()
        book.save(buffer)
        # A .xlsx named .xls (some tools do this), and a CSV named .xls.
        rows = import_services.read_file(
            SimpleUploadedFile("people.xls", buffer.getvalue()))
        self.assertEqual(rows[0]["employee_id"], "830011")
        rows = import_services.read_file(SimpleUploadedFile(
            "people.xls", b"ID,Name\r\n830011,Md. Hafizul Islam\r\n"))
        self.assertEqual(rows[0]["name"], "Md. Hafizul Islam")

    def test_a_web_page_saved_as_xls_says_so(self):
        upload = SimpleUploadedFile(
            "report.xls", b"<html><body><table><tr><td>ID</td></tr></table></body></html>")
        with self.assertRaisesMessage(ValidationError, "web page saved with a spreadsheet name"):
            import_services.read_file(upload)

    def test_another_ole_document_or_a_locked_workbook_is_refused(self):
        upload = SimpleUploadedFile("letter.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600)
        with self.assertRaisesMessage(ValidationError, "password-protected"):
            import_services.read_file(upload)


class VagueHeadingsTests(ImportCase):
    """A bare "ID" is often something else, sitting left of the real column.

    Accepting "ID" and "Code" made the first matching column win, so a row
    number called ID was imported as the employee ID - digits, so the preview
    accepted it, and ninety people arrived numbered 1 to 90 (found 2026-09-23).
    """

    def read(self, text):
        return import_services.read_file(
            SimpleUploadedFile("list.csv", text.encode("utf-8"), content_type="text/csv"))

    def test_a_row_number_called_id_does_not_beat_employee_id(self):
        rows = self.read("ID,Employee ID,Name\n1,445962,Ajay Ghosh\n2,445963,Dia Rahman\n")
        self.assertEqual([r["employee_id"] for r in rows], ["445962", "445963"])
        self.assertEqual([r["name"] for r in rows], ["Ajay Ghosh", "Dia Rahman"])

    def test_a_department_code_does_not_beat_employee_id(self):
        rows = self.read("Code,Employee ID,Name\n7,445962,Ajay Ghosh\n")
        self.assertEqual(rows[0]["employee_id"], "445962")

    def test_a_bare_id_is_still_used_when_it_is_the_only_one(self):
        """The client's own file: ID and Name, nothing else."""
        rows = self.read("ID,Name\n445962,Ajay Ghosh\n")
        self.assertEqual((rows[0]["employee_id"], rows[0]["name"]), ("445962", "Ajay Ghosh"))

    def test_employee_as_a_name_column_loses_to_name(self):
        rows = self.read("Employee ID,Employee,Name\n445962,Wrong,Ajay Ghosh\n")
        self.assertEqual(rows[0]["name"], "Ajay Ghosh")


def xlsx(*sheets, dimension=None):
    """An .xlsx with ``(title, rows)`` sheets. ``dimension`` overwrites the
    size each sheet claims, as some "export to Excel" programs get it wrong."""
    import re
    import zipfile

    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets:
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    if dimension is None:
        return buffer.getvalue()
    source, out = zipfile.ZipFile(buffer), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                data = re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="%s"' % dimension.encode(), data)
            target.writestr(item, data)
    return out.getvalue()


class RealFilesTests(ImportCase):
    """Files as clients actually send them (Ajay's upload refused on the live
    server, 2026-09-24): a title above the headings, a cover sheet, a size the
    file misreports, and a refusal nobody could misread."""

    def read(self, data, name="people.xlsx"):
        return import_services.read_file(SimpleUploadedFile(name, data))

    def test_a_title_above_the_headings_is_stepped_over(self):
        rows = self.read(xlsx(("Sheet1", [
            ["D Company Limited"],
            ["Employee list, September 2026"],
            [],
            ["SL", "Employee ID", "Name"],
            [1, 445962, "Ajay Kumar"],
            [2, 445963, "Dia Rahman"],
        ])))
        self.assertEqual([(r["line"], r["employee_id"], r["name"]) for r in rows],
                         [(5, "445962", "Ajay Kumar"), (6, "445963", "Dia Rahman")])

    def test_a_title_in_a_csv_is_stepped_over_too(self):
        data = "Employee list\n\nEmployee ID,Name\n445962,Ajay Kumar\n".encode()
        rows = import_services.read_file(SimpleUploadedFile("p.csv", data))
        self.assertEqual((rows[0]["line"], rows[0]["employee_id"]), (4, "445962"))

    def test_the_list_on_a_later_sheet_is_found(self):
        rows = self.read(xlsx(("Cover", [["Prepared by HR"]]),
                              ("Staff", [["ID", "Name"], [830011, "Md. Hafizul Islam"]])))
        self.assertEqual(rows[0]["employee_id"], "830011")

    def test_a_file_that_misreports_its_size_is_read_whole(self):
        # Without reset_dimensions, read-only openpyxl reads one cell of this.
        rows = self.read(xlsx(("Sheet1", [["Employee ID", "Name"], [445962, "Ajay Kumar"],
                                          [445963, "Dia Rahman"]]), dimension="A1:A1"))
        self.assertEqual([r["name"] for r in rows], ["Ajay Kumar", "Dia Rahman"])

    def test_more_spellings_and_bangla_headings(self):
        for headings in (("Emp. ID", "Employee's Name"), ("Staff ID", "Staff Name"),
                         ("ID No.", "Name of Employee"), ("কর্মচারী আইডি", "নাম")):
            with self.subTest(headings=headings):
                rows = import_services.read_file(
                    self.csv_file([["445962", "Ajay Kumar"]], headings=headings))
                self.assertEqual((rows[0]["employee_id"], rows[0]["name"]),
                                 ("445962", "Ajay Kumar"))

    def test_headings_too_far_down_are_not_guessed_at(self):
        filler = [[f"note {n}"] for n in range(import_services.HEADING_SEARCH_ROWS)]
        with self.assertRaisesMessage(ValidationError, "headings were not found"):
            self.read(xlsx(("Sheet1", filler + [["Employee ID", "Name"], [1, "A"]])))

    def test_the_refusal_says_what_it_wanted_where_it_looked_and_what_it_found(self):
        with self.assertRaises(ValidationError) as caught:
            self.read(xlsx(("Cover", [["Prepared by HR"]]),
                           ("Staff", [["Staff list"], ["SL", "Card No", "Staff Name"], [1, 7, "A"]])))
        message = " ".join(caught.exception.messages)
        self.assertIn("Wanted: a column headed Employee ID", message)
        self.assertIn("2 sheets (Cover, Staff)", message)
        self.assertIn('Closest was sheet "Staff", row 2. Found: SL, Card No, Staff Name.', message)
        self.assertIn("Missing: Employee ID.", message)

    def test_the_page_puts_the_refusal_at_the_top_whole(self):
        self.client.force_login(self.admin)
        page = self.client.post(self.url, {
            "branch": self.branch.pk,
            "upload": SimpleUploadedFile("staff.xlsx", xlsx(("Sheet1", [["SL", "Person"], [1, "A"]]))),
        })
        self.assertContains(page, "data-refusal")
        self.assertContains(page, "staff.xlsx was not imported.")
        self.assertContains(page, "Copy this whole box.")
        # Each part on its own line, and said once, not again under the field.
        self.assertContains(page, "<br>Looked in:")
        self.assertContains(page, "Wanted: a column headed Employee ID", count=1)


class ExistingPeopleTests(ImportCase):
    """A list already imported, with a few new people added, uploaded again
    (Nihal, 2026-09-24: 87 already in the software refused the 3 new ones).
    The people already here are skipped and left unchanged; the new ones import."""

    def setUp(self):
        super().setUp()
        with use_company(self.company):
            # Rahim, already an employee, under an Employee ID a file can carry.
            EmployeeAssignment.objects.filter(employee_code="E1").update(employee_code="770015")
        self.rahim = self.employee

    def test_the_new_people_import_and_the_existing_ones_are_skipped(self):
        page = self.upload([["770015", "Rahim"], ["7726978", "Nihal"], ["7726888", "Ajay"]])
        preview = page.context["preview"]
        self.assertEqual(preview["bad"], [])
        self.assertEqual(preview["counts"], {"total": 3, "bad": 0, "good": 2})
        self.assertEqual([r["employee_id"] for r in preview["existing"]], ["770015"])
        self.assertContains(page, "2 new people can be imported")
        self.assertContains(page, "1 already in the software is skipped and left unchanged.")
        self.assertContains(page, "Import 2 employees")

        before = self.count()
        page = self.confirm()
        self.assertEqual(self.count(), before + 2)
        self.assertContains(page, "2 employees imported")
        self.assertContains(page, "1 already in the software was skipped and left unchanged.")
        self.assertEqual(self.placement("7726978").employee.full_name, "Nihal")

    def test_the_same_name_leaves_them_unchanged(self):
        self.rahim.refresh_from_db()
        before = (self.rahim.first_name, self.rahim.last_name, self.rahim.updated_at)
        self.upload([["770015", self.rahim.full_name], ["7726978", "Nihal"]])
        self.confirm()
        self.rahim.refresh_from_db()
        self.assertEqual((self.rahim.first_name, self.rahim.last_name, self.rahim.updated_at), before)
        with use_company(self.company):
            self.assertEqual(EmployeeAssignment.objects.filter(employee_code="770015").count(), 1)

    def test_everyone_already_here_offers_nothing_to_import(self):
        page = self.upload([["770015", "Rahim"]])
        self.assertContains(page, "nobody new to import")
        self.assertNotContains(page, reverse("organization:employee_import_confirm"))
        with self.assertRaisesMessage(ValidationError, "already in the software"):
            import_services.commit(actor=self.admin, company_id=self.company.pk,
                                   rows=page.context["preview"]["existing"],
                                   branch_id=self.branch.pk)

    def test_a_real_mistake_still_stops_the_whole_file(self):
        page = self.upload([["770015", "Rahim"], ["7726978", "Nihal"], ["77A", "Bad Id"]])
        preview = page.context["preview"]
        self.assertEqual([r["line"] for r in preview["bad"]], [4])
        self.assertContains(page, "so the 1 new person is waiting too.")
        self.assertNotContains(page, "Import 1 employee")

    def test_a_person_new_at_the_preview_but_taken_by_confirm_refuses_the_lot(self):
        self.upload([["770015", "Rahim"], ["7726978", "Nihal"]])
        with use_company(self.company):
            EmployeeAssignment.objects.filter(employee=self.far).update(employee_code="7726978")
        before = self.count()
        self.assertContains(self.confirm(), "can no longer be imported")
        self.assertEqual(self.count(), before)

    def test_skipped_at_the_preview_stays_skipped(self):
        # The preview promised to leave 770015 alone; freeing that ID before
        # confirming must not turn it into a new employee nobody was shown.
        self.upload([["770015", "Rahim"], ["7726978", "Nihal"]])
        with use_company(self.company):
            EmployeeAssignment.objects.filter(employee_code="770015").update(employee_code="E1")
        before = self.count()
        self.confirm()
        self.assertEqual(self.count(), before + 1)
        with use_company(self.company):
            self.assertFalse(EmployeeAssignment.objects.filter(employee_code="770015").exists())

    def test_the_audit_names_who_was_skipped(self):
        self.upload([["770015", "Rahim"], ["7726978", "Nihal"]])
        self.confirm()
        entry = AuditLog.objects.get(action="employees.imported")
        self.assertEqual(entry.after_data["employee_codes"], ["7726978"])
        self.assertEqual(entry.after_data["skipped_codes"], ["770015"])

    def test_the_same_new_id_twice_in_the_file_is_still_a_mistake(self):
        page = self.upload([["7726978", "Nihal"], ["7726978", "Nihal Again"]])
        self.assertIn("also on row 2", page.context["preview"]["bad"][0]["errors"][0])


class RenameFromFileTests(ImportCase):
    """The same Employee ID with a different name in the file: the file's name
    replaces the old one (Nihal, 2026-09-24). Shown old and new before
    anything is written; audited; only for people the importer may edit."""

    def setUp(self):
        super().setUp()
        with use_company(self.company):
            EmployeeAssignment.objects.filter(employee_code="E1").update(employee_code="770015")
            self.rahim = self.employee
            self.rahim.middle_name = "Old"
            self.rahim.save(update_fields=["middle_name"])
        self.old_name = self.rahim.full_name

    def test_the_preview_shows_old_and_new_and_writes_nothing(self):
        page = self.upload([["770015", "Rahim Uddin Ahmed"], ["7726978", "Nihal"]])
        preview = page.context["preview"]
        self.assertEqual([(r["employee_id"], r["existing"], r["name"]) for r in preview["renames"]],
                         [("770015", self.old_name, "Rahim Uddin Ahmed")])
        self.assertEqual(preview["existing"], [])
        self.assertEqual(preview["counts"]["good"], 1)
        self.assertContains(page, "1 name will be updated to the file's.")
        self.assertContains(page, "Names to update (1)")
        self.assertContains(page, "Import 1 employee and update 1 name")
        self.rahim.refresh_from_db()
        self.assertEqual(self.rahim.full_name, self.old_name)

    def test_importing_replaces_the_whole_name(self):
        self.upload([["770015", "Rahim Uddin Ahmed"], ["7726978", "Nihal"]])
        page = self.confirm()
        self.rahim.refresh_from_db()
        self.assertEqual((self.rahim.first_name, self.rahim.middle_name, self.rahim.last_name),
                         ("Rahim Uddin", "", "Ahmed"))
        self.assertEqual(self.rahim.full_name, "Rahim Uddin Ahmed")
        self.assertEqual(self.rahim.updated_by, self.admin)
        self.assertContains(page, "1 employee imported")
        self.assertContains(page, "1 name updated from the file.")
        # Nothing but the name: same person, same placement, same Employee ID.
        placement = self.placement("770015")
        self.assertEqual(placement.employee_id, self.rahim.pk)
        self.assertEqual(placement.branch, self.branch)

    def test_a_file_that_only_renames_can_be_imported(self):
        page = self.upload([["770015", "Rahim Ahmed"]])
        self.assertContains(page, "Update 1 name")
        before = self.count()
        page = self.confirm()
        self.assertEqual(self.count(), before)
        self.assertContains(page, "Nobody new in people.csv. 1 name updated from the file.")
        self.rahim.refresh_from_db()
        self.assertEqual(self.rahim.full_name, "Rahim Ahmed")

    def test_a_change_of_capitals_or_spacing_counts_too(self):
        page = self.upload([["770015", self.old_name.upper()]])
        self.assertEqual(len(page.context["preview"]["renames"]), 1)
        page = self.upload([["770015", "  " + self.old_name.replace(" ", "   ") + " "]])
        # Only extra spaces: the same name once spaces are tidied, so skipped.
        self.assertEqual(page.context["preview"]["renames"], [])
        self.assertEqual(len(page.context["preview"]["existing"]), 1)

    def test_it_is_audited_with_the_old_and_new_name(self):
        self.upload([["770015", "Rahim Ahmed"], ["7726978", "Nihal"]])
        self.confirm()
        entry = AuditLog.objects.get(action="employees.imported")
        self.assertEqual(entry.after_data["renamed"], [{
            "employee_id": self.rahim.pk, "code": "770015",
            "from": self.old_name, "to": "Rahim Ahmed"}])

    def test_a_blank_name_leaves_them_as_they_are(self):
        page = self.upload([["770015", ""], ["7726978", "Nihal"]])
        preview = page.context["preview"]
        self.assertEqual(preview["bad"], [])
        self.assertEqual(preview["renames"], [])
        self.assertEqual([r["employee_id"] for r in preview["existing"]], ["770015"])

    def test_the_same_id_twice_with_two_names_is_a_mistake(self):
        page = self.upload([["770015", "Rahim Ahmed"], ["770015", "Rahim Khan"]])
        self.assertIn("also on row 2", page.context["preview"]["bad"][0]["errors"][0])
        self.assertNotContains(page, "Update 1 name")

    def test_a_rename_whose_id_is_gone_by_confirm_refuses_the_lot(self):
        self.upload([["770015", "Rahim Ahmed"], ["7726978", "Nihal"]])
        with use_company(self.company):
            EmployeeAssignment.objects.filter(employee_code="770015").update(employee_code="E1")
        before = self.count()
        self.assertContains(self.confirm(), "can no longer be imported")
        self.assertEqual(self.count(), before)
        self.rahim.refresh_from_db()
        self.assertEqual(self.rahim.full_name, self.old_name)

    def test_a_branch_manager_cannot_rename_or_see_someone_elsewhere(self):
        # Karim is in Chittagong; the manager runs Head Office only.
        with use_company(self.company):
            EmployeeAssignment.objects.filter(employee=self.far).update(employee_code="880001")
        page = self.upload([["880001", "Renamed By Manager"]], user=self.manager)
        row = page.context["preview"]["bad"][0]
        self.assertIn("outside your branches", row["errors"][0])
        self.assertNotContains(page, "Karim")
        self.assertEqual(page.context["preview"]["renames"], [])

    def test_a_branch_manager_renames_in_their_own_branch(self):
        self.upload([["770015", "Rahim Ahmed"]], user=self.manager)
        self.confirm()
        self.rahim.refresh_from_db()
        self.assertEqual(self.rahim.full_name, "Rahim Ahmed")
