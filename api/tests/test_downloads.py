"""The panel downloads the API first missed: one person's attendance calendar
and the Employees list (found by comparing every panel page with the API,
2026-10-05)."""

import io

from openpyxl import load_workbook

from api.tests.test_attendance import AttendanceApiTestCase


def sheet_text(response):
    book = load_workbook(io.BytesIO(response.content))
    return " ".join(str(cell.value) for row in book.active.iter_rows() for cell in row
                    if cell.value is not None)


class CalendarDownloadTests(AttendanceApiTestCase):
    def test_one_persons_month_as_a_pdf(self):
        got = self.api("GET", "/api/v1/attendance/calendar/export",
                       query=f"employee_id={self.rahim.pk}&year={self.late_day.year}"
                             f"&month={self.late_day.month}")
        self.assertEqual((got.status_code, got["Content-Type"]), (200, "application/pdf"))
        self.assertEqual(got.content[:5], b"%PDF-")
        missing = self.api("GET", "/api/v1/attendance/calendar/export")
        self.assertIn("employee_id", self.fields(missing))


class EmployeeDownloadTests(AttendanceApiTestCase):
    def test_the_list_as_excel_with_pay_for_those_who_see_it(self):
        got = self.api("GET", "/api/v1/employees/export")
        self.assertEqual(got.status_code, 200, got.content[:300])
        text = sheet_text(got)
        self.assertIn("Rahim", text)
        self.assertIn("Salary", text)
        pdf = self.api("GET", "/api/v1/employees/export", query="file_type=pdf")
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        wrong = self.api("GET", "/api/v1/employees/export", query="file_type=csv")
        self.assertEqual(wrong.status_code, 422)

    def test_a_key_without_payroll_read_gets_no_pay(self):
        key = self.key("employees:read")
        got = self.as_key(key, "GET", "/api/v1/employees/export")
        self.assertEqual(got.status_code, 200, got.content[:300])
        text = sheet_text(got)
        self.assertIn("Rahim", text)
        self.assertNotIn("Salary", text)
        self.assertNotIn("30000", text)
