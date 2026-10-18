"""Shared set-up for the inventory tests: a company with the module on, its
owner, and a few masters. Not a test module itself."""

import datetime
from decimal import Decimal

from django.contrib.auth import get_user_model

from accounts.models import CompanyMembership
from common.tenant import use_company
from inventory import masters
from inventory.models import Item, ItemCategory, ItemKind, Uom, Warehouse
from organization.models import Branch
from tenants import platform_services
from tenants.models import Feature

User = get_user_model()
D = Decimal
DAY = datetime.date(2026, 10, 1)


def filled_template(filename):
    """The template ``filename`` with its Example rows typed into the sheets
    to fill, from row 2 - what a client sends back. An uploaded file."""
    import io
    from pathlib import Path

    from django.conf import settings
    from django.core.files.uploadedfile import SimpleUploadedFile
    from openpyxl import load_workbook

    book = load_workbook(Path(settings.BASE_DIR) / "docs" / "inventory" / "templates" / filename)
    fill = [name for name in book.sheetnames if name not in ("How to fill", "Example")]
    blocks, current = {}, None
    for row in book["Example"].iter_rows(values_only=True):
        if isinstance(row[0], str) and row[0].endswith(" (example)"):
            current = row[0][: -len(" (example)")]
            continue
        if any(value is not None for value in row):
            blocks.setdefault(current or fill[0], []).append(row)
    for name, rows in blocks.items():
        for r, values in enumerate(rows[1:], start=2):
            for c, value in enumerate(values, start=1):
                book[name].cell(row=r, column=c, value=value)
    buffer = io.BytesIO()
    book.save(buffer)
    return SimpleUploadedFile(filename, buffer.getvalue())


class InventoryFixture:
    """Mix into a TestCase: ``self.company``, ``self.owner``, ``self.branch``,
    units ``self.kg`` / ``self.pcs``, stores ``self.rms`` / ``self.fgw``,
    items ``self.sugar`` (batches + expiry, KG) and ``self.bearing`` (PCS)."""

    password = "Pw-81723!inv"

    def setUp(self):
        super().setUp()
        self.root = User.objects.create_superuser(email="root@inv.test", password=self.password)
        self.company = platform_services.create_platform_company(
            actor=self.root, values={"name": "Inventory Foods Ltd"})
        self.enable_module(True)
        self.owner = self.member("owner@inv.test", "owner")
        with use_company(self.company):
            self.branch = Branch.objects.get(is_default=True)
        masters.add_standard_units(actor=self.owner, company_id=self.company.pk)
        with use_company(self.company):
            self.kg = Uom.objects.get(code="KG")
            self.pcs = Uom.objects.get(code="PCS")
            self.bag = Uom.objects.get(code="BAG")
        self.raw = self.category("SWT", "Sweetener", ItemKind.RAW)
        self.spares = self.category("BRG", "Bearing", ItemKind.SPARE)
        self.rms = self.store("RMS", "Raw material store", Warehouse.Kind.RAW)
        self.fgw = self.store("FGW", "Finished goods", Warehouse.Kind.FINISHED)
        self.sugar = self.item("RM-0001", "Sugar", ItemKind.RAW, self.raw, self.kg,
                               track_batch=True, track_expiry=True, shelf_life_days=365)
        self.bearing = self.item("SP-0001", "Bearing 6205", ItemKind.SPARE, self.spares, self.pcs)

    # --- helpers ---------------------------------------------------------

    def enable_module(self, on):
        feature = Feature.objects.get(code="inventory")
        platform_services.set_company_feature(
            actor=self.root, company_id=self.company.pk, feature_id=feature.pk,
            effect="enable" if on else "disable", reason="Inventory tests")

    def member(self, email, role):
        user = User.objects.create_user(email=email, password=self.password)
        CompanyMembership.all_objects.create(company=self.company, user=user, role=role,
                                             status="active")
        return user

    def category(self, code, name, kind):
        return masters.create(masters.CATEGORY, actor=self.owner, company_id=self.company.pk,
                              values={"code": code, "name": name, "kind": kind, "parent": None})

    def store(self, code, name, kind, branch=None):
        return masters.create(masters.WAREHOUSE, actor=self.owner, company_id=self.company.pk,
                              values={"branch": branch or self.branch, "code": code, "name": name,
                                      "kind": kind, "notes": ""})

    def item(self, code, name, kind, category, uom, **extra):
        values = {"code": code, "name": name, "kind": kind, "category": category,
                  "base_uom": uom, **extra}
        return masters.create(masters.ITEM, actor=self.owner, company_id=self.company.pk,
                              values=values)

    def tenant(self):
        return use_company(self.company)

    def reload(self, obj):
        return type(obj).all_objects.get(pk=obj.pk)

    def get_item(self, pk):
        return Item.all_objects.select_related("base_uom").get(pk=pk)

    def get_category(self, code):
        return ItemCategory.all_objects.get(company=self.company, code=code)
