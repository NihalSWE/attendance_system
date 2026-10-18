"""Phase 10: going live - the readiness checklist, count sheets that import
back as opening stock, labels for the opening stock, and the parallel run."""

import datetime
import io
from io import StringIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from common.tenant import use_company
from inventory import golive, imports, masters, opening
from inventory.models import BarcodeLabel, OpeningStock
from inventory.testing import D, InventoryFixture

DAY = datetime.date(2026, 11, 1)


class GoLiveFixture(InventoryFixture):
    def post_opening(self):
        kw = {"actor": self.owner, "company_id": self.company.pk}
        doc = opening.create_document(**kw, values={"work_date": DAY, "note": ""})
        opening.add_line(**kw, document_id=doc.pk, values={
            "warehouse": self.rms, "location": None, "item": self.sugar, "batch_no": "S1", "supplier_batch_no": "",
            "mfg_date": None, "expiry_date": datetime.date(2027, 9, 30), "qty": D("50"), "uom": self.kg,
            "unit_cost": D("80"), "note": ""})
        opening.add_line(**kw, document_id=doc.pk, values={
            "warehouse": self.rms, "location": None, "item": self.bearing, "batch_no": "", "supplier_batch_no": "",
            "mfg_date": None, "expiry_date": None, "qty": D("4"), "uom": self.pcs, "unit_cost": D("300"),
            "note": ""})
        opening.post_document(**kw, document_id=doc.pk)
        return doc

    def statuses(self):
        return {check.title: check.status for _title, checks in golive.readiness(self.company.pk) for check in checks}


class ReadinessTests(GoLiveFixture, TestCase):
    def test_the_checklist_moves_from_not_ready_to_ready(self):
        found = self.statuses()
        self.assertEqual(found["Inventory settings are not saved yet"], golive.FAIL)
        self.assertEqual(found["0 opening stock document(s) posted"], golive.FAIL)
        self.assertEqual(found["Nobody has inventory.issue yet"], golive.WARN)
        with self.assertRaises(CommandError):
            call_command("golive_check", company=self.company.pk, stdout=StringIO())

        masters.update_settings(actor=self.owner, company_id=self.company.pk, values={
            "valuation_scope": "company", "expiry_alert_days": 30, "slow_moving_days": 90, "alert_emails": "",
            "small_label": "50x25", "large_label": "100x50", "gs1_labels": False})
        doc = self.post_opening()
        found = self.statuses()
        self.assertEqual(found["Inventory settings saved"], golive.OK)
        self.assertEqual(found["1 opening stock document(s) posted"], golive.OK)
        self.assertEqual(found["1 batch(es) in stock have no printed label"], golive.WARN)
        self.assertEqual(found["The ledger and the balances agree"], golive.OK)
        out = StringIO()
        call_command("golive_check", company=self.company.pk, stdout=out)
        self.assertIn("0 to fix", out.getvalue())

        # Printing the opening stock's labels clears the label warning.
        made = golive.opening_labels(actor=self.owner, company_id=self.company.pk, document_id=doc.pk)
        self.assertEqual(sorted(label.kind for label in made), ["batch", "item"])
        from inventory import labels

        labels.print_labels(actor=self.owner, company_id=self.company.pk, label_ids=[l.pk for l in made])
        self.assertNotIn("1 batch(es) in stock have no printed label", self.statuses())


class CountSheetTests(GoLiveFixture, TestCase):
    def test_a_filled_count_sheet_imports_as_opening_stock(self):
        content, name = golive.count_sheet(self.company.pk, self.rms)
        self.assertEqual(name, "count-sheet-RMS.xlsx")
        book = load_workbook(io.BytesIO(content))
        sheet = book.active
        heading = next(row[0].row for row in sheet.iter_rows() if row[0].value == "Store code")
        columns = {cell.value: cell.column for cell in sheet[heading]}
        listed = {sheet.cell(r, columns["Item code"]).value: r for r in range(heading + 1, sheet.max_row + 1)}
        self.assertEqual(set(listed), {"RM-0001", "SP-0001"})      # no usual stores yet: every item
        # The counter writes sugar's quantity, batch and expiry; the bearing is not found.
        row = listed["RM-0001"]
        sheet.cell(row, columns["Quantity"]).value = 48
        sheet.cell(row, columns["Batch no."]).value = "S1"
        sheet.cell(row, columns["Expiry date"]).value = datetime.date(2027, 9, 30)
        sheet.cell(row, columns["Unit cost (BDT)"]).value = 80
        filled = io.BytesIO()
        book.save(filled)
        upload = SimpleUploadedFile("count.xlsx", filled.getvalue())
        data = imports.read("opening", upload)
        rows, _extra = imports.check("opening", self.owner, self.company.pk, data, DAY)
        self.assertEqual(sorted((r.label, r.action) for r in rows),
                         [("RMS · RM-0001 S1", "add"), ("RMS · SP-0001", "skip")])
        document = imports.commit("opening", actor=self.owner, company_id=self.company.pk, data=data, work_date=DAY)
        with use_company(self.company):
            [line] = document.lines.all()
        self.assertEqual((line.item, line.qty, line.batch_no), (self.sugar, D("48"), "S1"))

    def test_the_pdf_and_the_page(self):
        content, name = golive.count_sheet(self.company.pk, self.rms, fmt="pdf")
        self.assertTrue(content.startswith(b"%PDF"))
        self.client.force_login(self.owner)
        page = self.client.get(reverse("inventory:golive"))
        self.assertContains(page, "Count sheets")
        self.assertContains(page, 'data-menu="inventory"')
        sheet = self.client.get(reverse("inventory:golive_count_sheet"), {"store": self.rms.pk, "file_type": "xlsx"})
        self.assertIn("spreadsheetml", sheet["Content-Type"])


class ParallelRunTests(GoLiveFixture, TestCase):
    def test_the_old_systems_list_against_the_software(self):
        self.post_opening()
        old = SimpleUploadedFile("old.csv", b"Item code,Store code,Quantity\nRM-0001,RMS,48\nSP-0001,RMS,4\n"
                                            b"XX-9,RMS,1\n")
        differences, problems = golive.compare(self.company.pk, old)
        self.assertEqual([(d.item, d.old, d.ours, d.status) for d in differences],
                         [("RM-0001", D("48"), D("50"), "Different"), ("SP-0001", D("4"), D("4"), "Same")])
        self.assertEqual(problems, ["Row 3: item XX-9 is not in the software."])
        self.client.force_login(self.owner)
        page = self.client.post(reverse("inventory:golive_compare"), {"upload": SimpleUploadedFile(
            "old.csv", b"Item code,Quantity\nRM-0001,45\n")})
        self.assertContains(page, "2 of 2 differ")     # the bearing is only in the software
        file = self.client.post(reverse("inventory:golive_compare"), {"download": "on", "upload": SimpleUploadedFile(
            "old.csv", b"Item code,Quantity\nRM-0001,45\n")})
        self.assertIn("spreadsheetml", file["Content-Type"])

    def test_opening_labels_print_from_the_page(self):
        doc = self.post_opening()
        self.client.force_login(self.owner)
        self.assertContains(self.client.get(reverse("inventory:opening_detail", args=[doc.pk])), "Print all labels")
        pdf = self.client.post(reverse("inventory:opening_labels", args=[doc.pk]), {"copies": "1"})
        self.assertEqual(pdf["Content-Type"], "application/pdf")
        self.assertEqual(BarcodeLabel.all_objects.filter(printed_count=1).count(), 2)
        self.assertEqual(OpeningStock.all_objects.get().status, "posted")
